#!/usr/bin/env python3
"""Convert EE .doc/.docx test case documents to structured JSON."""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

AZURE_REPO_URL = "https://aspentech-alm.visualstudio.com/AspenTech/_git/EE_Main"

SECTION_RE = re.compile(
    r"^(?P<section>\d+(?:\.\d+)*)\s+(?P<title>.+?)\s*\((?P<id>\d+)\)\s*$"
)
SUBCASE_RE = re.compile(r"^(?P<id>[a-zA-Z]\))\s*(?P<title>.+)$")
STEP_RE = re.compile(r"^(?P<id>\d+\))\s*(?P<text>.+)$")


@dataclass
class Subcase:
    id: str
    title: str
    steps: list[str] = field(default_factory=list)


@dataclass
class CaseDocument:
    id: str
    section: str
    title: str
    category: str
    tags: list[str]
    steps: list[str] = field(default_factory=list)
    subcases: list[Subcase] = field(default_factory=list)


def main() -> None:
    args = parse_args()
    workspace = Path(args.workspace).resolve()
    source_repo = workspace / "EE_Main"
    output_dir = Path(args.output).resolve()

    if not source_repo.exists() or args.refresh:
        clone_repo(source_repo, args.pat)

    output_dir.mkdir(parents=True, exist_ok=True)
    converted = 0

    for doc_path in iter_docs(source_repo):
        text_lines = read_document_lines(doc_path)
        parsed_cases = parse_cases(text_lines)
        if not parsed_cases:
            continue

        target_file = output_dir / f"{doc_path.stem}.json"
        target_file.write_text(
            json.dumps([to_dict(case) for case in parsed_cases], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        converted += 1

    print(f"Converted {converted} documents to JSON in {output_dir}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert EE docs to JSON")
    parser.add_argument(
        "--workspace",
        default="/tmp/ee_docs",
        help="Workspace used for cloning Azure DevOps repository",
    )
    parser.add_argument(
        "--output",
        default="src/data",
        help="Directory to write converted JSON files",
    )
    parser.add_argument(
        "--refresh",
        action="store_true",
        help="Force a fresh clone of the source repository",
    )
    parser.add_argument(
        "--pat",
        default=os.getenv("AZURE_DEVOPS_PAT", ""),
        help="Azure DevOps PAT (defaults to AZURE_DEVOPS_PAT env var)",
    )
    return parser.parse_args()


def clone_repo(dest: Path, pat: str) -> None:
    if not pat:
        raise SystemExit(
            "AZURE_DEVOPS_PAT is required to clone the private source repository."
        )

    if dest.exists():
        shutil.rmtree(dest)

    dest.parent.mkdir(parents=True, exist_ok=True)
    clone_url = AZURE_REPO_URL.replace("https://", f"https://:{pat}@")
    subprocess.run(["git", "clone", clone_url, str(dest)], check=True)


def iter_docs(root: Path) -> Iterable[Path]:
    for extension in ("*.docx", "*.doc"):
        yield from root.rglob(extension)


def read_document_lines(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        from docx import Document

        document = Document(path)
        return [normalize_line(paragraph.text) for paragraph in document.paragraphs]

    return [normalize_line(line) for line in extract_doc_text(path).splitlines()]


def extract_doc_text(path: Path) -> str:
    antiword = shutil.which("antiword")
    if antiword:
        completed = subprocess.run(
            [antiword, str(path)],
            check=False,
            capture_output=True,
            text=True,
        )
        if completed.returncode == 0 and completed.stdout.strip():
            return completed.stdout

    try:
        import textract  # type: ignore

        extracted = textract.process(str(path))
        return extracted.decode("utf-8", errors="ignore")
    except ImportError as exc:  # pragma: no cover - runtime dependent
        raise RuntimeError(
            "Unable to parse .doc. Install antiword or textract to enable fallback parsing."
        ) from exc


def parse_cases(lines: list[str]) -> list[CaseDocument]:
    cases: list[CaseDocument] = []
    current_case: CaseDocument | None = None
    current_subcase: Subcase | None = None

    for raw in lines:
        line = raw.strip()
        if not line:
            continue

        section_match = SECTION_RE.match(line)
        if section_match:
            current_case = CaseDocument(
                id=section_match.group("id"),
                section=section_match.group("section"),
                title=section_match.group("title"),
                category=infer_category(section_match.group("title")),
                tags=infer_tags(section_match.group("title")),
            )
            cases.append(current_case)
            current_subcase = None
            continue

        if current_case is None:
            continue

        subcase_match = SUBCASE_RE.match(line)
        if subcase_match:
            normalized_subcase_id = subcase_match.group("id").replace(")", ""
            ).lower()
            section_id = current_case.section
            current_subcase = Subcase(
                id=f"{section_id}{normalized_subcase_id}",
                title=subcase_match.group("title"),
            )
            current_case.subcases.append(current_subcase)
            continue

        step_match = STEP_RE.match(line)
        step_text = step_match.group("text") if step_match else line

        if current_subcase is not None:
            current_subcase.steps.append(step_text)
        else:
            current_case.steps.append(step_text)

    return [case for case in cases if case.subcases or case.steps]


def infer_category(title: str) -> str:
    first_word = title.split()[0] if title.split() else "General"
    return first_word.capitalize()


def infer_tags(title: str) -> list[str]:
    words = re.findall(r"[A-Za-z0-9+.]+", title)
    tags = [word for word in words if len(word) > 2]
    return list(dict.fromkeys(tags))[:6]


def normalize_line(line: str) -> str:
    return re.sub(r"\s+", " ", line).strip()


def to_dict(case: CaseDocument) -> dict[str, object]:
    return {
        "id": case.id,
        "section": case.section,
        "title": case.title,
        "category": case.category,
        "tags": case.tags,
        "steps": case.steps,
        "subcases": [
            {"id": subcase.id, "title": subcase.title, "steps": subcase.steps}
            for subcase in case.subcases
        ],
    }


if __name__ == "__main__":
    main()
